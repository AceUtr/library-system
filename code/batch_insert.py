"""
批量将 .c 源文件代码插入到 Word 文档中
"""
import zipfile, os, shutil, re
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ===== 第一步: 修复 docx 中的损坏引用 =====

doc_path = r"c:\Users\wang\Desktop\程序设计训练作业\答辩\图书管理系统设计说明书_中式版.docx"
temp_dir = os.environ['TEMP'] + '\\docx_fix'
fixed_path = os.environ['TEMP'] + '\\target_fixed.docx'

if os.path.exists(temp_dir):
    shutil.rmtree(temp_dir, ignore_errors=True)

print("Step 1: 修复文档中的损坏引用...")
with zipfile.ZipFile(doc_path, 'r') as z:
    z.extractall(temp_dir)

rels_path = os.path.join(temp_dir, 'word', '_rels', 'document.xml.rels')
if os.path.exists(rels_path):
    with open(rels_path, 'r', encoding='utf-8') as f:
        content = f.read()
    fixed = re.sub(r'<Relationship[^>]*Target="[^"]*NULL[^"]*"[^>]*/>', '', content)
    fixed = re.sub(r"<Relationship[^>]*Target='[^']*NULL[^']*'[^>]*/>", '', fixed)
    if fixed != content:
        print("  -> 已移除 NULL 引用")
        with open(rels_path, 'w', encoding='utf-8') as f:
            f.write(fixed)

if os.path.exists(fixed_path):
    os.remove(fixed_path)

with zipfile.ZipFile(fixed_path, 'w', zipfile.ZIP_DEFLATED) as new_z:
    for root, dirs, files in os.walk(temp_dir):
        for file in files:
            full_path = os.path.join(root, file)
            arc_name = os.path.relpath(full_path, temp_dir).replace('\\', '/')
            new_z.write(full_path, arc_name)

shutil.rmtree(temp_dir, ignore_errors=True)
print("  -> 修复完成")

# ===== 第二步: 插入源代码 =====

print("\nStep 2: 插入源代码...")
doc = Document(fixed_path)
print(f"  文档已打开 ({len(doc.paragraphs)} 个段落)")

def add_section_title(doc, title):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(title)
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(12)
    run.bold = True
    p_format = paragraph.paragraph_format
    p_format.space_before = Pt(12)
    p_format.space_after = Pt(6)

def add_code_block(doc, content, font_name='Consolas', font_size=10.5, fill_color='F5F5F5'):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(content)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(font_size)

    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:val'), 'clear')
    shading_elm.set(qn('w:color'), 'auto')
    shading_elm.set(qn('w:fill'), fill_color)
    paragraph._p.get_or_add_pPr().append(shading_elm)

    p_format = paragraph.paragraph_format
    p_format.line_spacing = Pt(13)
    p_format.space_before = Pt(0)
    p_format.space_after = Pt(0)

source_dir = r"c:\Users\wang\Desktop\程序设计训练作业\Library system"
c_files = [
    ("main.c",   "main.c — 主程序入口"),
    ("borrow.c", "borrow.c — 借阅管理模块"),
    ("purchase.c", "purchase.c — 图书采购管理模块"),
    ("fileio.c", "fileio.c — 文件读写模块"),
    ("stats.c",  "stats.c — 统计分析模块"),
    ("ui.c",     "ui.c — 命令行用户界面"),
]

add_section_title(doc, "附录：系统源代码")

for filename, title in c_files:
    filepath = os.path.join(source_dir, filename)
    if not os.path.exists(filepath):
        print(f"  [跳过] 文件不存在: {filepath}")
        continue
    with open(filepath, 'r', encoding='utf-8') as f:
        code = f.read()
    print(f"  插入: {filename} ({len(code)} 字符)")
    add_section_title(doc, title)
    add_code_block(doc, code)

# 原文档被 Word 锁定，保存到同目录下新文件
output_path = r"c:\Users\wang\Desktop\程序设计训练作业\答辩\图书管理系统设计说明书_含源代码.docx"
doc.save(output_path)
print(f"\n完成! 已保存至: {output_path}")
